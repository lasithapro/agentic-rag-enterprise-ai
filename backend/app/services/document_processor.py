"""Document processing service for PDFs, Word, and text files."""
import os
import uuid
from pathlib import Path
from typing import List, Tuple, Optional
import aiofiles
from fastapi import UploadFile
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from app.core.config import get_settings
from app.core.logging_config import get_logger

logger = get_logger(__name__)
settings = get_settings()


class DocumentProcessor:
    """Processes documents from various sources into text chunks."""

    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            length_function=len,
            add_start_index=True,
        )
        self.upload_dir = Path(settings.UPLOAD_DIR)
        self.upload_dir.mkdir(parents=True, exist_ok=True)

    async def save_upload_file(self, file: UploadFile) -> Tuple[str, Path]:
        """Save uploaded file and return document_id and path."""
        document_id = str(uuid.uuid4())
        file_extension = Path(file.filename).suffix.lower()
        file_path = self.upload_dir / f"{document_id}{file_extension}"

        async with aiofiles.open(file_path, "wb") as f:
            content = await file.read()
            await f.write(content)

        logger.info("Saved upload file", document_id=document_id, filename=file.filename)
        return document_id, file_path

    def process_pdf(self, file_path: Path, document_id: str, filename: str) -> List[Document]:
        """Process PDF file into document chunks."""
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(file_path))
            documents = []
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text.strip():
                    doc = Document(
                        page_content=text,
                        metadata={
                            "document_id": document_id,
                            "filename": filename,
                            "page_number": page_num + 1,
                            "source": str(file_path),
                            "document_type": "pdf",
                        },
                    )
                    documents.append(doc)
            chunks = self.text_splitter.split_documents(documents)
            logger.info("Processed PDF", document_id=document_id, chunks=len(chunks))
            return chunks
        except Exception as e:
            logger.error("Failed to process PDF", document_id=document_id, error=str(e))
            raise

    def process_word(self, file_path: Path, document_id: str, filename: str) -> List[Document]:
        """Process Word document into document chunks."""
        try:
            from docx import Document as DocxDocument
            doc = DocxDocument(str(file_path))
            full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            document = Document(
                page_content=full_text,
                metadata={
                    "document_id": document_id,
                    "filename": filename,
                    "source": str(file_path),
                    "document_type": "word",
                },
            )
            chunks = self.text_splitter.split_documents([document])
            logger.info("Processed Word document", document_id=document_id, chunks=len(chunks))
            return chunks
        except Exception as e:
            logger.error("Failed to process Word document", document_id=document_id, error=str(e))
            raise

    def process_text(self, file_path: Path, document_id: str, filename: str) -> List[Document]:
        """Process plain text file into document chunks."""
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
            document = Document(
                page_content=text,
                metadata={
                    "document_id": document_id,
                    "filename": filename,
                    "source": str(file_path),
                    "document_type": "text",
                },
            )
            chunks = self.text_splitter.split_documents([document])
            logger.info("Processed text file", document_id=document_id, chunks=len(chunks))
            return chunks
        except Exception as e:
            logger.error("Failed to process text file", document_id=document_id, error=str(e))
            raise

    async def process_web_url(self, url: str, document_id: str) -> List[Document]:
        """Process web URL content into document chunks."""
        try:
            import httpx
            from bs4 import BeautifulSoup

            async with httpx.AsyncClient(timeout=30.0) as client:
                response = await client.get(url, follow_redirects=True)
                response.raise_for_status()

            soup = BeautifulSoup(response.text, "html.parser")
            for script in soup(["script", "style", "nav", "footer", "header"]):
                script.decompose()

            text = soup.get_text(separator="\n", strip=True)
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            text = "\n".join(lines)

            document = Document(
                page_content=text,
                metadata={
                    "document_id": document_id,
                    "url": url,
                    "source": url,
                    "document_type": "web",
                },
            )
            chunks = self.text_splitter.split_documents([document])
            logger.info("Processed web URL", document_id=document_id, url=url, chunks=len(chunks))
            return chunks
        except Exception as e:
            logger.error("Failed to process web URL", document_id=document_id, url=url, error=str(e))
            raise

    def process_file(self, file_path: Path, document_id: str, filename: str) -> List[Document]:
        """Route file to appropriate processor based on extension."""
        extension = file_path.suffix.lower()
        if extension == ".pdf":
            return self.process_pdf(file_path, document_id, filename)
        elif extension in [".doc", ".docx"]:
            return self.process_word(file_path, document_id, filename)
        elif extension in [".txt", ".md"]:
            return self.process_text(file_path, document_id, filename)
        else:
            raise ValueError(f"Unsupported file type: {extension}")
